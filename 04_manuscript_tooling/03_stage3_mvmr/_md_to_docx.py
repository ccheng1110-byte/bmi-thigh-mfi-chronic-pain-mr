# -*- coding: utf-8 -*-
"""
P4b: Convert the Markdown files in the submission package to Word (.docx) and export main Tables 1-6 to .xlsx.
Requires: python-docx, openpyxl (venv: C:/Users/青涯/.workbuddy/binaries/python/envs/default)
"""
import re
import sys
import io
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

OUT = Path(r"d:/SX/TMFI_CPI_MR/05_submission_JCSM_2026-09-20")
TARGETS = [
    OUT / "01_Main_text" / "Main_Text_JCSM_2026-09-20.md",
    OUT / "02_Tables" / "Main_Tables_1-6_combined.md",
    OUT / "04_Supplementary" / "Supplementary_Material_JCSM_2026-09-20.md",
    OUT / "05_Cover_letter" / "Cover_letter_JCSM_2026-09-20.md",
    OUT / "06_Reporting_checklist" / "STROBE-MR_Checklist_TMFI_BMI_CPI_2026-09-20.md",
    OUT / "07_Originality_package" / "原创性包_最近邻比较与主贡献声明_2026-09-20.md",
]

INLINE = re.compile(
    r"(\*\*.+?\*\*|`[^`]+`|<sub>.*?</sub>|<sup>.*?</sup>|\[[^\]]+\]\([^)]+\)|\*[^*\n]+?\*)"
)


def add_rich(p, text, base_bold=False, base_italic=False):
    """Render markdown inline markup as docx runs."""
    pos = 0
    for m in INLINE.finditer(text):
        if m.start() > pos:
            p.add_run(text[pos:m.start()])
        tok = m.group(0)
        r = None
        if tok.startswith("**") and tok.endswith("**"):
            r = p.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith("`") and tok.endswith("`"):
            r = p.add_run(tok[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(10)
        elif tok.startswith("<sub>"):
            r = p.add_run(tok[5:-6]); r.font.subscript = True
        elif tok.startswith("<sup>"):
            r = p.add_run(tok[5:-6]); r.font.superscript = True
        elif tok.startswith("["):
            mm = re.match(r"\[([^\]]+)\]\(([^)]+)\)", tok)
            r = p.add_run(mm.group(1)); r.font.color.rgb = RGBColor(0x06, 0x3B, 0x9E); r.underline = True
        elif tok.startswith("*") and tok.endswith("*"):
            r = p.add_run(tok[1:-1]); r.italic = True
        if r is None:
            r = p.add_run(tok)
        if base_bold:
            r.bold = True
        if base_italic:
            r.italic = True
        pos = m.end()
    if pos < len(text):
        r = p.add_run(text[pos:])
        r.bold = base_bold
        r.italic = base_italic


def split_row(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def md_to_docx(md_path: Path) -> Path:
    text = md_path.read_text(encoding="utf-8").replace("\r\n", "\n")
    lines = text.split("\n")

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Pt(72)

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()

        # tables
        if stripped.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:\-\|]+\|$", lines[i + 1].strip()):
            header = split_row(stripped)
            i += 2
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i].strip()))
                i += 1
            tbl = doc.add_table(rows=1, cols=len(header))
            tbl.style = "Table Grid"
            for j, h in enumerate(header):
                cell = tbl.rows[0].cells[j]
                cell.text = ""
                add_rich(cell.paragraphs[0], h, base_bold=True)
            for row in rows:
                cells = tbl.add_row().cells
                for j in range(min(len(row), len(header))):
                    cells[j].text = ""
                    add_rich(cells[j].paragraphs[0], row[j])
            doc.add_paragraph()
            continue

        if not stripped:
            i += 1
            continue

        # horizontal rule
        if re.match(r"^-{3,}$", stripped):
            p = doc.add_paragraph()
            i += 1
            continue

        # headings
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            lvl = len(m.group(1))
            p = doc.add_heading(level=min(lvl, 4))
            p.text = ""
            add_rich(p, m.group(2))
            i += 1
            continue

        # lists
        m = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", line)
        if m:
            style = "List Bullet" if m.group(2) in ("-", "*") else "List Number"
            p = doc.add_paragraph(style=style)
            add_rich(p, m.group(3))
            i += 1
            continue

        # blockquotes
        if stripped.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(24)
            add_rich(p, stripped.lstrip("> ").strip(), base_italic=True)
            i += 1
            continue

        # ordinary paragraphs (soft line breaks merged)
        buf = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if (not nxt) or nxt.startswith("#") or nxt.startswith("|") or nxt.startswith(">"):
                break
            if re.match(r"^(\s*)([-*]|\d+\.)\s+", lines[i]):
                break
            buf.append(nxt)
            i += 1
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_rich(p, " ".join(buf))
    docx_path = md_path.with_suffix(".docx")
    doc.save(docx_path)
    return docx_path


def tables_to_xlsx(md_path: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side

    text = md_path.read_text(encoding="utf-8").replace("\r\n", "\n")
    lines = text.split("\n")
    wb = Workbook()
    wb.remove(wb.active)
    thin = Border(*[Side(style="thin", color="BFBFBF")] * 4)
    cur = None
    title = "Table"
    i = 0
    blocks = []
    while i < len(lines):
        s = lines[i].strip()
        m = re.match(r"^### (Table \d+)\.\s*(.*)$", s)
        if m:
            title = m.group(1) + " — " + m.group(2)[:40]
            i += 1
            continue
        if s.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:\-\|]+\|$", lines[i + 1].strip()):
            header = split_row(s)
            i += 2
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i].strip()))
                i += 1
            blocks.append((title, header, rows))
            continue
        i += 1

    def clean(c):
        c = re.sub(r"<sub>|</sub>|<sup>|</sup>", "", c)
        c = c.replace("**", "").replace("`", "")
        c = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", c)
        return c.strip()

    for idx, (t, header, rows) in enumerate(blocks, start=1):
        ws = wb.create_sheet(t.split(" — ")[0].replace(" ", ""))
        ws["A1"] = t
        ws["A1"].font = Font(bold=True, size=12)
        for j, h in enumerate(header, start=1):
            c = ws.cell(row=2, column=j, value=clean(h))
            c.font = Font(bold=True)
            c.border = thin
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for r, row in enumerate(rows, start=3):
            for j in range(1, min(len(row), len(header)) + 1):
                c = ws.cell(row=r, column=j, value=clean(row[j - 1]))
                c.border = thin
                c.alignment = Alignment(wrap_text=True, vertical="top")
        for j, h in enumerate(header, start=1):
            ws.column_dimensions[ws.cell(row=2, column=j).column_letter].width = \
                min(max(14, len(clean(h)) + 6), 46)
    xlsx = md_path.parent / "Main_Tables_1-6.xlsx"
    wb.save(xlsx)
    return xlsx


if __name__ == "__main__":
    for t in TARGETS:
        if not t.exists():
            print("MISSING:", t)
            continue
        d = md_to_docx(t)
        print("docx:", d.name, d.stat().st_size, "bytes")
    x = tables_to_xlsx(OUT / "02_Tables" / "Main_Tables_1-6_combined.md")
    print("xlsx:", x.name, x.stat().st_size, "bytes")
